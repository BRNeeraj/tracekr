import io
import os
import re
import argparse
from datetime import date, datetime
from functools import wraps
from pathlib import Path
from urllib.parse import quote_plus

import qrcode
from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text, UniqueConstraint, func
from sqlalchemy.exc import SQLAlchemyError
from werkzeug.security import check_password_hash, generate_password_hash

load_dotenv()

db = SQLAlchemy()


def database_uri():
    """Build the production MySQL URL from Render environment variables."""
    explicit_url = os.getenv("DATABASE_URL", "").strip()
    if explicit_url:
        return explicit_url.replace("mysql://", "mysql+pymysql://", 1)

    db_host = os.getenv("DB_HOST", "").strip()
    db_name = os.getenv("DB_NAME", "").strip()
    if db_host and db_name:
        db_port = os.getenv("DB_PORT", "3306").strip()
        db_user = quote_plus(os.getenv("DB_USER", "").strip())
        db_password = quote_plus(os.getenv("DB_PASSWORD", ""))
        return f"mysql+pymysql://{db_user}:{db_password}@{db_host}:{db_port}/{db_name}?charset=utf8mb4"

    if os.getenv("APP_ENV", "development").lower() == "production":
        raise RuntimeError("Production requires DATABASE_URL or DB_HOST, DB_PORT, DB_USER, DB_PASSWORD, and DB_NAME.")
    return "sqlite:///hackathon_tracker.db"


class Participant(db.Model):
    __tablename__ = "participants"
    id = db.Column(db.Integer, primary_key=True)
    participant_id = db.Column(db.String(30), unique=True, nullable=False, index=True)
    name = db.Column(db.String(120), nullable=False)
    phone = db.Column(db.String(30), nullable=False)
    team_name = db.Column(db.String(120), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    records = db.relationship("ServiceRecord", back_populates="participant", cascade="all, delete-orphan")


class Service(db.Model):
    __tablename__ = "services"
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), unique=True, nullable=False)
    is_active = db.Column(db.Boolean, default=True, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    records = db.relationship("ServiceRecord", back_populates="service", cascade="all, delete-orphan")


class ServiceRecord(db.Model):
    __tablename__ = "service_records"
    id = db.Column(db.Integer, primary_key=True)
    participant_id = db.Column(db.Integer, db.ForeignKey("participants.id", ondelete="CASCADE"), nullable=False)
    service_id = db.Column(db.Integer, db.ForeignKey("services.id", ondelete="CASCADE"), nullable=False)
    event_date = db.Column(db.Date, default=date.today, nullable=False)
    scanned_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    participant = db.relationship("Participant", back_populates="records")
    service = db.relationship("Service", back_populates="records")
    __table_args__ = (UniqueConstraint("participant_id", "service_id", "event_date", name="uq_service_per_day"),)


class Admin(db.Model):
    __tablename__ = "admins"
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)


def create_app(test_config=None):
    if os.getenv("APP_ENV", "development").lower() == "production":
        if not os.getenv("SECRET_KEY") or os.getenv("SECRET_KEY") == "change-this-secret-in-production":
            raise RuntimeError("Production requires a non-default SECRET_KEY environment variable.")
        if not os.getenv("ADMIN_PASSWORD"):
            raise RuntimeError("Production requires an ADMIN_PASSWORD environment variable.")
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.getenv("SECRET_KEY", "change-this-secret-in-production"),
        SQLALCHEMY_DATABASE_URI=database_uri(),
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        SQLALCHEMY_ENGINE_OPTIONS={"pool_pre_ping": True, "pool_recycle": 280},
        EVENT_NAME=os.getenv("EVENT_NAME", "Hackathon 2026"),
    )
    if test_config:
        app.config.update(test_config)
    db.init_app(app)

    with app.app_context():
        try:
            db.create_all()
            if not Admin.query.filter_by(username=os.getenv("ADMIN_USERNAME", "admin")).first():
                db.session.add(Admin(username=os.getenv("ADMIN_USERNAME", "admin"), password_hash=generate_password_hash(os.getenv("ADMIN_PASSWORD", "admin123"))))
                db.session.commit()
            if Service.query.count() == 0:
                db.session.add_all([Service(name=n) for n in ["Breakfast", "Morning Tea", "Coffee", "Lunch", "Snacks", "Evening Tea", "Dinner"]])
                db.session.commit()
        except SQLAlchemyError:
            db.session.rollback()
            app.logger.exception("Database initialization failed; the app will remain available for health diagnostics.")

    def login_required(fn):
        @wraps(fn)
        def wrapped(*args, **kwargs):
            if "admin_id" not in session:
                return redirect(url_for("login", next=request.path))
            return fn(*args, **kwargs)
        return wrapped

    @app.context_processor
    def globals_for_templates():
        return {"event_name": app.config["EVENT_NAME"], "today": date.today()}

    @app.errorhandler(500)
    def server_error(error):
        db.session.rollback()
        return render_template("error.html", message="The database or server is temporarily unavailable."), 500

    @app.get("/health")
    def health():
        try:
            db.session.execute(text("SELECT 1"))
            return {"status": "ok", "database": "ok"}
        except SQLAlchemyError:
            db.session.rollback()
            return {"status": "degraded", "database": "unavailable"}, 503

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            admin = Admin.query.filter_by(username=request.form.get("username", "").strip()).first()
            if admin and check_password_hash(admin.password_hash, request.form.get("password", "")):
                session["admin_id"] = admin.id
                session["username"] = admin.username
                return redirect(request.args.get("next") or url_for("dashboard"))
            flash("Invalid username or password.", "error")
        return render_template("login.html")

    @app.get("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.get("/")
    @login_required
    def dashboard():
        participants = Participant.query.count()
        services = Service.query.order_by(Service.name).all()
        stats = []
        for service in services:
            taken = ServiceRecord.query.filter_by(service_id=service.id, event_date=date.today()).count()
            stats.append({"service": service, "taken": taken, "total": participants, "percent": round(taken / participants * 100) if participants else 0})
        return render_template("dashboard.html", participants=participants, stats=stats)

    def next_participant_id():
        latest = Participant.query.order_by(Participant.id.desc()).first()
        return f"HACK-{(latest.id + 1 if latest else 1):04d}"

    @app.route("/participants", methods=["GET", "POST"])
    @login_required
    def participants():
        if request.method == "POST":
            name, phone, team = [request.form.get(k, "").strip() for k in ("name", "phone", "team_name")]
            if not name or not team or not re.fullmatch(r"[0-9+() .-]{7,20}", phone):
                flash("Enter a name, team, and valid phone number.", "error")
            else:
                participant = Participant(participant_id=next_participant_id(), name=name, phone=phone, team_name=team)
                db.session.add(participant)
                db.session.commit()
                flash(f"Participant {participant.participant_id} created.", "success")
                return redirect(url_for("participant_view", participant_id=participant.participant_id))
        query = request.args.get("q", "").strip()
        items = Participant.query
        if query:
            like = f"%{query}%"
            items = items.filter(db.or_(Participant.name.ilike(like), Participant.participant_id.ilike(like), Participant.team_name.ilike(like)))
        return render_template("participants.html", participants=items.order_by(Participant.id.desc()).all(), query=query)

    @app.route("/participants/<participant_id>/edit", methods=["GET", "POST"])
    @login_required
    def participant_edit(participant_id):
        participant = Participant.query.filter_by(participant_id=participant_id).first_or_404()
        if request.method == "POST":
            name, phone, team = [request.form.get(k, "").strip() for k in ("name", "phone", "team_name")]
            if not name or not team or not re.fullmatch(r"[0-9+() .-]{7,20}", phone):
                flash("Enter a name, team, and valid phone number.", "error")
            else:
                participant.name, participant.phone, participant.team_name = name, phone, team
                db.session.commit()
                flash("Participant updated.", "success")
                return redirect(url_for("participant_view", participant_id=participant_id))
        return render_template("participant_form.html", participant=participant)

    @app.post("/participants/<participant_id>/delete")
    @login_required
    def participant_delete(participant_id):
        participant = Participant.query.filter_by(participant_id=participant_id).first_or_404()
        db.session.delete(participant)
        db.session.commit()
        flash("Participant deleted.", "success")
        return redirect(url_for("participants"))

    @app.get("/participants/<participant_id>")
    @login_required
    def participant_view(participant_id):
        return render_template("participant_view.html", participant=Participant.query.filter_by(participant_id=participant_id).first_or_404())

    @app.get("/participants/<participant_id>/qr")
    @login_required
    def participant_qr(participant_id):
        participant = Participant.query.filter_by(participant_id=participant_id).first_or_404()
        image = qrcode.make(participant.participant_id)
        output = io.BytesIO()
        image.save(output, "PNG")
        output.seek(0)
        return send_file(output, mimetype="image/png", download_name=f"{participant.participant_id}.png")

    @app.get("/print-qrs")
    @login_required
    def print_qrs():
        return render_template("print_qr.html", participants=Participant.query.order_by(Participant.id).all())

    @app.route("/services", methods=["GET", "POST"])
    @login_required
    def services():
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            if not name:
                flash("Service name is required.", "error")
            elif Service.query.filter(func.lower(Service.name) == name.lower()).first():
                flash("That service already exists.", "error")
            else:
                db.session.add(Service(name=name))
                db.session.commit()
                flash("Service added.", "success")
            return redirect(url_for("services"))
        return render_template("services.html", services=Service.query.order_by(Service.created_at.desc()).all())

    @app.post("/services/<int:service_id>/toggle")
    @login_required
    def service_toggle(service_id):
        service = Service.query.get_or_404(service_id)
        service.is_active = not service.is_active
        db.session.commit()
        return redirect(url_for("services"))

    @app.get("/scanner")
    @login_required
    def scanner():
        return render_template("scanner.html", services=Service.query.filter_by(is_active=True).order_by(Service.name).all())

    @app.post("/api/scan")
    @login_required
    def scan():
        participant = Participant.query.filter_by(participant_id=request.json.get("participant_id", "").strip()).first()
        service = Service.query.get(request.json.get("service_id"))
        if not participant:
            return {"status": "invalid", "message": "Participant ID not found."}, 404
        if not service or not service.is_active:
            return {"status": "error", "message": "Select an active service."}, 400
        existing = ServiceRecord.query.filter_by(participant_id=participant.id, service_id=service.id, event_date=date.today()).first()
        if existing:
            return {"status": "duplicate", "participant": participant.name, "team": participant.team_name, "service": service.name, "time": existing.scanned_at.strftime("%I:%M %p").lstrip("0")}
        record = ServiceRecord(participant_id=participant.id, service_id=service.id, event_date=date.today())
        db.session.add(record)
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            existing = ServiceRecord.query.filter_by(participant_id=participant.id, service_id=service.id, event_date=date.today()).first()
            if existing:
                return {"status": "duplicate", "participant": participant.name, "team": participant.team_name, "service": service.name, "time": existing.scanned_at.strftime("%I:%M %p").lstrip("0")}
            return {"status": "error", "message": "Could not save this scan."}, 500
        return {"status": "success", "participant": participant.name, "team": participant.team_name, "service": service.name, "time": record.scanned_at.strftime("%I:%M %p").lstrip("0")}

    @app.get("/reports")
    @login_required
    def reports():
        service_id = request.args.get("service", type=int)
        selected = Service.query.get(service_id) if service_id else Service.query.filter_by(is_active=True).order_by(Service.name).first()
        status = request.args.get("status", "all")
        query = request.args.get("q", "").strip()
        rows = []
        if selected:
            records = {r.participant_id: r for r in ServiceRecord.query.filter_by(service_id=selected.id, event_date=date.today()).all()}
            for p in Participant.query.order_by(Participant.participant_id).all():
                record = records.get(p.id)
                if query and query.lower() not in f"{p.name} {p.participant_id} {p.team_name}".lower():
                    continue
                if status == "taken" and not record or status == "not_taken" and record:
                    continue
                rows.append((p, record))
        taken = ServiceRecord.query.filter_by(service_id=selected.id, event_date=date.today()).count() if selected else 0
        return render_template("reports.html", services=Service.query.order_by(Service.name).all(), selected=selected, rows=rows, taken=taken, total_participants=Participant.query.count(), query=query, status=status)

    @app.get("/export.xlsx")
    @login_required
    def export_xlsx():
        wb = Workbook()
        ws = wb.active
        ws.title = "Participants"
        participants = Participant.query.order_by(Participant.participant_id).all()
        write_sheet(ws, ["Participant ID", "Name", "Phone", "Team"], [[p.participant_id, p.name, p.phone, p.team_name] for p in participants])
        records_ws = wb.create_sheet("Service Records")
        records = ServiceRecord.query.join(Participant).join(Service).filter(ServiceRecord.event_date == date.today()).order_by(ServiceRecord.scanned_at).all()
        write_sheet(records_ws, ["Participant ID", "Name", "Phone", "Team", "Service", "Date", "Time", "Status"], [[r.participant.participant_id, r.participant.name, r.participant.phone, r.participant.team_name, r.service.name, r.event_date.isoformat(), r.scanned_at.strftime("%I:%M %p"), "Taken"] for r in records])
        summary = wb.create_sheet("Summary")
        write_sheet(summary, ["Service", "Total Participants", "Taken", "Not Taken"], [[s.name, len(participants), ServiceRecord.query.filter_by(service_id=s.id, event_date=date.today()).count(), len(participants) - ServiceRecord.query.filter_by(service_id=s.id, event_date=date.today()).count()] for s in Service.query.order_by(Service.name).all()])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"hackathon-report-{date.today().isoformat()}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    return app


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for row in rows:
        ws.append(row)
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 35)
        ws.column_dimensions[column[0].column_letter].width = width


app = create_app()

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Run the Hackathon Tracker server")
    parser.add_argument("--https", action="store_true", help="serve HTTPS using certs/dev-cert.pem and certs/dev-key.pem")
    args = parser.parse_args()
    use_https = args.https or os.getenv("FLASK_HTTPS", "0") == "1"
    ssl_context = None
    if use_https:
        cert_file = Path(os.getenv("SSL_CERT_FILE", "certs/dev-cert.pem"))
        key_file = Path(os.getenv("SSL_KEY_FILE", "certs/dev-key.pem"))
        if not cert_file.exists() or not key_file.exists():
            raise SystemExit("HTTPS certificate files are missing. Run: python scripts/generate_dev_cert.py")
        ssl_context = (str(cert_file), str(key_file))
    app.run(
        host=os.getenv("FLASK_HOST", "0.0.0.0"),
        port=int(os.getenv("PORT", os.getenv("FLASK_PORT", "5000"))),
        debug=os.getenv("FLASK_DEBUG", "0") == "1",
        ssl_context=ssl_context,
    )
